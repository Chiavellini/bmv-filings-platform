"""
test_segments_helpers.py — metric-list → Segments helpers (offline, no rendering).

These helpers were lifted out of the former Streamlit dashboard into
``src.excel.segments_sheet`` so the CLI (scripts/build_segments.py) and tests can
reuse them. Pure functions only: the fuzzy-query → SegmentRow mapping, the title
helper, the xlsx-bytes builder, and outline-spec detection.
"""

from __future__ import annotations

from io import BytesIO

import openpyxl
import pandas as pd

from src.excel.segments_sheet import (
    SegmentRow,
    build_segments_xlsx,
    load_metric_defs,
    segments_from_queries,
    segments_title,
)


def test_segments_from_queries_resolves_dedupes_and_drops_unmatched():
    defs = load_metric_defs()
    segs = segments_from_queries(
        ["revenue", "ebitda", "revenue", "definitely not a metric xyz"], defs)
    keys = [s.key for s in segs]
    assert "revenue" in keys
    assert "ebitda" in keys
    assert keys.count("revenue") == 1            # de-duplicated
    assert len(keys) == 2                          # unmatched dropped
    assert all(s.label for s in segs)              # labels populated


def test_segments_from_queries_empty():
    assert segments_from_queries([], load_metric_defs()) == []


def test_segments_title_with_and_without_ticker():
    assert segments_title("Grupo Bimbo", "BIMBOA") == "BIMBOA: Grupo Bimbo"
    assert segments_title("Acme", "") == "Acme"


def test_build_segments_xlsx_returns_valid_workbook():
    rows = []
    for year, base in ((2023, 100), (2024, 110)):
        for q in (1, 2, 3, 4):
            rows.append({"period": f"{year}-{q}T", "revenue": base + q, "ebitda": base // 2})
    df = pd.DataFrame(rows)
    segs = [SegmentRow("Revenue", "revenue"), SegmentRow("EBITDA", "ebitda")]

    data = build_segments_xlsx("ACME: Test Co", segs, df)
    assert isinstance(data, bytes) and len(data) > 0

    wb = openpyxl.load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Segments"]
    ws = wb["Segments"]
    assert ws["A2"].value == "ACME: Test Co"
    assert ws["H5"].value == "1Q24A"          # current year auto-detected
    assert ws["B11"].value == "Revenue"        # first segment data row


def test_company_outline_spec_detected_for_walmex():
    # The real walmex.yaml resolves to a structured outline spec.
    from pathlib import Path

    from src.model.financial_model import load_config
    from src.excel.segments_sheet import _load_segments_spec

    root = Path(__file__).resolve().parents[1]
    cfg_path = root / "configs" / "walmex.yaml"
    cfg = load_config(str(cfg_path))
    spec = _load_segments_spec(cfg, str(cfg_path))
    assert spec and spec.get("outline")
    assert "Revenues" in spec.get("sections", [])
