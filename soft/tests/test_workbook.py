"""Phase-4 tests — the valuation workbook renders with live formulas (synthetic fixtures)."""
from __future__ import annotations

from openpyxl import load_workbook

from src.bloomberg.schema import BloombergPack
from src.coverage.fundamentals import Fundamentals
from src.coverage.spec import CoverageSpec, Labelled, Peer
from src.coverage.valuation import build_model
from src.sheets.valuation_sheet import build_valuation_workbook


def _fixtures():
    fund = Fundamentals(
        slug="walmex", frame=None, periods=["2024-4T"], current_period="2024-4T",
        ltm={"revenue": 800_000, "ebitda": 80_000, "net_income": 40_000,
             "operating_income": 60_000, "gross_profit": 180_000, "total_assets": 500_000,
             "inventory": 70_000, "cfo": 60_000, "capex": 20_000,
             "ebitda_cam": 10_000, "sales_floor_mexico": 6_000_000, "total_stores": 3_800},
        annual={2023: {"ebitda": 75_000, "net_income": 37_000, "revenue": 750_000},
                2024: {"ebitda": 80_000, "net_income": 40_000, "revenue": 800_000}},
        flow_keys={"revenue", "ebitda", "net_income", "operating_income", "gross_profit",
                   "cfo", "capex", "ebitda_cam"})
    pack = BloombergPack(
        slug="walmex",
        subject={"px_last": 60.0, "shares_out": 17_400, "net_debt": 50_000,
                 "minority_interest": 0.0, "eps_ntm": 2.6, "dvd_yield": 1.6,
                 "total_equity": 250_000, "fcf_ltm": 40_000, "total_assets": 500_000,
                 "replacement_cost_per_sqm": 12_000},
        peers={"chedraui": {"px_last": 120.0, "shares_out": 960, "net_debt": 20_000,
                            "sales_ltm": 260_000, "ebitda_ltm": 22_000, "net_income_ltm": 8_000,
                            "total_equity": 60_000, "fcf_ltm": 6_000}},
        segment_multiples={"mexico": 9.5, "cam": 7.0},
        history={2023: 72.0, 2024: 62.0},
        macro={"gdp_growth": 1.4})
    spec = CoverageSpec(
        slug="walmex", name="Walmex", ticker="WALMEX* MM",
        blocks=["snapshot_multiples", "historical_multiples", "sum_of_the_parts",
                "replacement_value", "financial_analysis", "macro_sector"],
        peers=[Peer("chedraui", "CHDRAUI", "CHDRAUI")],
        segments=[Labelled("mexico", "Mexico"), Labelled("cam", "Central America")],
        macro=[Labelled("gdp_growth", "Mexico GDP growth, %")])
    return spec, fund, pack


def test_workbook_renders_with_live_formulas(tmp_path):
    spec, fund, pack = _fixtures()
    model = build_model(spec, fund, pack)
    out = build_valuation_workbook(model, spec, fund, pack, tmp_path / "Walmex.xlsx")
    assert out.exists()

    wb = load_workbook(out)  # keep formulas as strings
    assert "Valuation" in wb.sheetnames
    ws = wb["Valuation"]

    formulas = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    # market cap + EV + 6 subject multiples + peer medians → several live formulas
    assert any("*" in f for f in formulas)       # market cap = price*shares
    assert any(f.startswith("=MEDIAN(") for f in formulas)  # peer-median cells
    assert len(formulas) >= 8

    # the title cell names the company
    assert "Walmex" in str(ws["A1"].value)
