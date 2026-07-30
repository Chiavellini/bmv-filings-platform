"""Tests for the strict Excel-formula evaluator (src.coverage.formula_eval)."""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from src.coverage.formula_eval import (
    DivisionByZero,
    FormulaError,
    evaluate_worksheet,
    make_evaluator,
)


def _ws(cells: dict):
    wb = Workbook()
    ws = wb.active
    for coord, val in cells.items():
        ws[coord] = val
    return ws


def test_arithmetic_and_refs():
    ws = _ws({"B2": 60.0, "B3": 100.0, "B4": "=B2*B3", "B5": "=B4/B2", "B6": "=B4+B2-B3"})
    ev = make_evaluator(ws)
    assert ev.cell("B4") == pytest.approx(6000.0)
    assert ev.cell("B5") == pytest.approx(100.0)
    assert ev.cell("B6") == pytest.approx(6000.0 + 60.0 - 100.0)


def test_precedence():
    ws = _ws({"A1": 2.0, "A2": 3.0, "A3": 4.0, "A4": "=A1+A2*A3"})
    assert make_evaluator(ws).cell("A4") == pytest.approx(2 + 3 * 4)


def test_median():
    ws = _ws({"C1": 10.0, "D1": 20.0, "E1": 30.0, "F1": "=MEDIAN(C1,D1,E1)"})
    assert make_evaluator(ws).cell("F1") == pytest.approx(20.0)


def test_division_by_zero_raises():
    ws = _ws({"B2": 5.0, "B3": 0.0, "B4": "=B2/B3"})
    with pytest.raises(DivisionByZero):
        make_evaluator(ws).cell("B4")


def test_unknown_function_raises():
    ws = _ws({"B2": 1.0, "B3": "=SUMX(B2)"})
    with pytest.raises(FormulaError):
        make_evaluator(ws).cell("B3")


def test_circular_reference_raises():
    ws = _ws({"B2": "=B3", "B3": "=B2"})
    with pytest.raises(FormulaError):
        make_evaluator(ws).cell("B2")


def test_evaluate_worksheet_collects_formulas():
    ws = _ws({"B2": 2.0, "B3": 3.0, "B4": "=B2*B3"})
    values, formulas = evaluate_worksheet(ws)
    assert formulas == {"B4": "=B2*B3"}
    assert values["B4"] == pytest.approx(6.0)
