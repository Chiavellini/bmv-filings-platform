"""Analyst-provided metric-sheet ingestion and fidelity checks.

The onboarding Markdown remains the executable specification because it pins
canonical extractor keys, but it must not drift from the analyst's original
CSV/XLSX request.  This module reads two deliberately small contracts:

* normalized tables with ``section`` and ``label`` columns (``key`` optional);
* legacy matrix sheets where a company name/ticker marks the start of a block
  and the requested labels are in the first non-empty cell to its left.

It does not guess metric keys or classify rows.  That judgment stays explicit
in ``inputs/<slug>.md`` and is checked by the onboarding compiler.
"""

from __future__ import annotations

from dataclasses import dataclass
import csv
import hashlib
from pathlib import Path
import re
import unicodedata


class AnalystMetricSheetError(ValueError):
    """The analyst sheet cannot be read or does not contain the requested block."""


@dataclass(frozen=True, slots=True)
class AnalystMetricSheet:
    """Ordered labels and immutable source identity for one analyst request."""

    path: Path
    company: str
    labels: tuple[str, ...]
    keys: tuple[str | None, ...]
    kinds: tuple[str | None, ...]
    sha256: str
    sheet_name: str | None = None


@dataclass(frozen=True, slots=True)
class MetricFidelityResult:
    """Sequence comparison between the analyst sheet and executable outline."""

    matches: bool
    expected: tuple[str, ...]
    actual: tuple[str, ...]
    first_difference: str | None = None


@dataclass(frozen=True, slots=True)
class MetricKeyFidelityResult:
    """Comparison of any canonical keys explicitly supplied by the analyst."""

    matches: bool
    first_difference: str | None = None


@dataclass(frozen=True, slots=True)
class MetricKindFidelityResult:
    """Comparison of analyst row roles (section versus requested row)."""

    matches: bool
    first_difference: str | None = None


_GENERIC_HEADERS = re.compile(
    r"^(?:segments?|segment data)(?:\s*\([^)]*\)|\s+in\s+.+)?$",
    re.IGNORECASE,
)


def normalize_label(value: object) -> str:
    """Normalize display-only differences without hiding semantic drift."""

    text = unicodedata.normalize("NFKC", str(value or ""))
    text = text.replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip().casefold()


def normalize_company(value: object) -> str:
    """Loose identity normalization for company/ticker block markers."""

    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text.casefold())


def load_analyst_metric_sheet(
    path: str | Path,
    company: str,
    *,
    aliases: tuple[str, ...] = (),
    sheet_name: str | None = None,
) -> AnalystMetricSheet:
    """Read one company block from an analyst CSV or XLSX workbook."""

    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise AnalystMetricSheetError(f"analyst metric sheet not found: {source}")
    suffix = source.suffix.lower()
    try:
        if suffix == ".csv":
            with source.open(encoding="utf-8-sig", newline="") as fh:
                rows = [list(row) for row in csv.reader(fh)]
            selected_sheet = None
        elif suffix in {".xlsx", ".xlsm"}:
            rows, selected_sheet = _read_xlsx_rows(
                source,
                company=company,
                aliases=aliases,
                requested_sheet=sheet_name,
            )
        else:
            raise AnalystMetricSheetError(
                f"unsupported analyst metric sheet format {suffix!r}; use .csv or .xlsx"
            )
    except AnalystMetricSheetError:
        raise
    except Exception as exc:  # malformed ZIP/XML/encoding is an intake defect
        raise AnalystMetricSheetError(
            f"cannot read analyst metric sheet {source.name}: "
            f"{type(exc).__name__}: {exc}"
        ) from exc

    labels, keys, kinds = _contract_from_rows(rows, company=company, aliases=aliases)
    if not labels:
        raise AnalystMetricSheetError(
            f"no metric labels found for {company!r} in {source}"
        )
    return AnalystMetricSheet(
        path=source,
        company=company,
        labels=tuple(labels),
        keys=tuple(keys),
        kinds=tuple(kinds),
        sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        sheet_name=selected_sheet,
    )


def compare_metric_sequences(
    expected: tuple[str, ...] | list[str],
    actual: tuple[str, ...] | list[str],
) -> MetricFidelityResult:
    """Compare ordered section/row labels and identify the first mismatch."""

    exp = tuple(str(value).strip() for value in expected if str(value).strip())
    got = tuple(str(value).strip() for value in actual if str(value).strip())
    exp_norm = tuple(normalize_label(value) for value in exp)
    got_norm = tuple(normalize_label(value) for value in got)
    if exp_norm == got_norm:
        return MetricFidelityResult(True, exp, got)

    shared = min(len(exp), len(got))
    index = next(
        (i for i in range(shared) if exp_norm[i] != got_norm[i]),
        shared,
    )
    expected_value = exp[index] if index < len(exp) else "<end of analyst sheet>"
    actual_value = got[index] if index < len(got) else "<end of onboarding outline>"
    return MetricFidelityResult(
        False,
        exp,
        got,
        f"row {index + 1}: analyst={expected_value!r}, outline={actual_value!r}",
    )


def compare_metric_keys(
    labels: tuple[str, ...] | list[str],
    expected: tuple[str | None, ...] | list[str | None],
    actual: tuple[str | None, ...] | list[str | None],
) -> MetricKeyFidelityResult:
    """Enforce analyst-declared canonical keys while allowing blank key cells.

    Legacy analyst sheets usually contain display labels only.  A normalized
    sheet may additionally provide a ``key`` column; every non-empty key in that
    column is authoritative and must match the corresponding Markdown pin.
    """

    names = tuple(str(value).strip() for value in labels)
    exp = tuple(
        str(value).strip() if value is not None and str(value).strip() else None
        for value in expected
    )
    got = tuple(
        str(value).strip() if value is not None and str(value).strip() else None
        for value in actual
    )
    if len(exp) != len(names) or len(got) != len(names):
        return MetricKeyFidelityResult(
            False,
            "analyst labels, analyst keys, and onboarding keys have different lengths",
        )
    for index, (expected_key, actual_key) in enumerate(zip(exp, got)):
        if expected_key is None:
            continue
        if expected_key.casefold() != (actual_key or "").casefold():
            return MetricKeyFidelityResult(
                False,
                f"row {index + 1} ({names[index]!r}): "
                f"analyst key={expected_key!r}, outline key={actual_key!r}",
            )
    return MetricKeyFidelityResult(True)


def compare_metric_kinds(
    labels: tuple[str, ...] | list[str],
    expected: tuple[str | None, ...] | list[str | None],
    actual: tuple[str | None, ...] | list[str | None],
) -> MetricKindFidelityResult:
    """Reject reclassification of a requested row as a section (or vice versa)."""

    names = tuple(str(value).strip() for value in labels)
    exp = tuple(value.casefold() if isinstance(value, str) and value else None
                for value in expected)
    got = tuple(value.casefold() if isinstance(value, str) and value else None
                for value in actual)
    if len(exp) != len(names) or len(got) != len(names):
        return MetricKindFidelityResult(
            False,
            "analyst labels, analyst row kinds, and onboarding row kinds have different lengths",
        )
    for index, (expected_kind, actual_kind) in enumerate(zip(exp, got)):
        if expected_kind is None:
            continue
        if expected_kind != actual_kind:
            return MetricKindFidelityResult(
                False,
                f"row {index + 1} ({names[index]!r}): "
                f"analyst kind={expected_kind!r}, outline kind={actual_kind!r}",
            )
    return MetricKindFidelityResult(True)


def _read_xlsx_rows(
    path: Path,
    *,
    company: str,
    aliases: tuple[str, ...],
    requested_sheet: str | None,
) -> tuple[list[list[object]], str]:
    """Read values only; workbook authoring remains in the Excel builder."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - declared project dependency
        raise AnalystMetricSheetError("openpyxl is required to read .xlsx metric sheets") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if requested_sheet:
            if requested_sheet not in workbook.sheetnames:
                raise AnalystMetricSheetError(
                    f"worksheet {requested_sheet!r} not found in {path.name}"
                )
            candidates = [workbook[requested_sheet]]
        else:
            identities = {
                normalize_company(value)
                for value in (company, *aliases)
                if normalize_company(value)
            }
            named = [
                worksheet
                for worksheet in workbook.worksheets
                if normalize_company(worksheet.title) in identities
            ]
            candidates = named + [
                worksheet for worksheet in workbook.worksheets if worksheet not in named
            ]

        last_error: AnalystMetricSheetError | None = None
        for worksheet in candidates:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            try:
                _labels_from_rows(rows, company=company, aliases=aliases)
            except AnalystMetricSheetError as exc:
                last_error = exc
                continue
            return rows, worksheet.title
        if last_error is not None:
            raise last_error
        raise AnalystMetricSheetError(f"workbook has no worksheets: {path}")
    finally:
        workbook.close()


def _labels_from_rows(
    rows: list[list[object]],
    *,
    company: str,
    aliases: tuple[str, ...],
) -> list[str]:
    """Extract normalized-table or legacy block labels from raw cell rows."""

    return _contract_from_rows(rows, company=company, aliases=aliases)[0]


def _contract_from_rows(
    rows: list[list[object]],
    *,
    company: str,
    aliases: tuple[str, ...],
) -> tuple[list[str], list[str | None], list[str | None]]:
    """Extract ordered display labels and optional analyst-declared keys."""

    if not rows:
        raise AnalystMetricSheetError("analyst metric sheet is empty")
    normalized = _normalized_table_contract(rows)
    if normalized is not None:
        return normalized
    labels = _legacy_block_labels(rows, company=company, aliases=aliases)
    return labels, [None] * len(labels), [None] * len(labels)


def _normalized_table_labels(rows: list[list[object]]) -> list[str] | None:
    contract = _normalized_table_contract(rows)
    return contract[0] if contract is not None else None


def _normalized_table_contract(
    rows: list[list[object]],
) -> tuple[list[str], list[str | None], list[str | None]] | None:
    header_index = next((i for i, row in enumerate(rows) if any(row)), None)
    if header_index is None:
        return [], [], []
    header = [normalize_label(value) for value in rows[header_index]]
    if "section" not in header or "label" not in header:
        return None
    sec_col, label_col = header.index("section"), header.index("label")
    key_col = header.index("key") if "key" in header else None
    labels: list[str] = []
    keys: list[str | None] = []
    kinds: list[str | None] = []
    current_section = ""
    for row in rows[header_index + 1 :]:
        section = _cell_text(row, sec_col)
        label = _cell_text(row, label_col)
        if section and normalize_label(section) != normalize_label(current_section):
            labels.append(section)
            keys.append(None)
            kinds.append("section")
            current_section = section
        if label:
            labels.append(label)
            key = _cell_text(row, key_col) if key_col is not None else ""
            keys.append(key or None)
            kinds.append("row")
    return labels, keys, kinds


def _legacy_block_labels(
    rows: list[list[object]],
    *,
    company: str,
    aliases: tuple[str, ...],
) -> list[str]:
    identities = {
        normalize_company(value)
        for value in (company, *aliases)
        if normalize_company(value)
    }
    marker: tuple[int, int] | None = None
    discovered: list[str] = []
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            text = str(value or "").strip()
            if not text:
                continue
            if column_index > 0:
                discovered.append(text)
            if normalize_company(text) in identities:
                marker = (row_index, column_index)
                break
        if marker is not None:
            break
    if marker is None:
        hint = ", ".join(dict.fromkeys(discovered[:8])) or "none"
        raise AnalystMetricSheetError(
            f"company block {company!r} not found; observed marker values: {hint}"
        )

    start, marker_col = marker
    labels: list[str] = []
    for row_index in range(start, len(rows)):
        row = rows[row_index]
        if row_index > start and _cell_text(row, marker_col):
            break
        candidates = [
            str(value).strip()
            for value in row[:marker_col]
            if str(value or "").strip()
        ]
        if not candidates:
            continue
        label = candidates[0]
        if not labels and _GENERIC_HEADERS.fullmatch(label):
            continue
        labels.append(label)
    return labels


def _cell_text(row: list[object], index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


__all__ = [
    "AnalystMetricSheet",
    "AnalystMetricSheetError",
    "MetricFidelityResult",
    "MetricKeyFidelityResult",
    "MetricKindFidelityResult",
    "compare_metric_keys",
    "compare_metric_kinds",
    "compare_metric_sequences",
    "load_analyst_metric_sheet",
    "normalize_company",
    "normalize_label",
]
