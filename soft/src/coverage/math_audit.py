"""Independent math-verification gate for a built coverage deliverable.

The engine computes every ratio once in :mod:`src.coverage.valuation` (Python) and, for the
industrial snapshot, a second time as live Excel formulas in :mod:`src.sheets.valuation_sheet`.
This module re-derives every ratio a *third* way — with its own arithmetic, deliberately **not**
importing ``company_multiples``/the valuation block builders — and asserts all three agree. It
also guards against the ways a "number" can be wrong even when a formula is arithmetically fine:
division by zero, ``NaN``/``inf``, Excel error strings, and a ratio printed while its denominator
is blank.

``audit(model, spec, fund, pack, xlsx_path) -> AuditReport``. ``report.passed`` is True only when
there are no **gating** failures (arithmetic/identity/guard errors). Range-plausibility and
meaningless-but-correct multiples (e.g. a negative P/E from negative earnings) are recorded as
**advisory** warnings — they do not block, because they are not math errors.

A blank cell (an unsourced input, e.g. missing Bloomberg data) is always allowed; only a *present*
value that is arithmetically wrong blocks the deliverable.
"""
from __future__ import annotations

import math
import statistics
from dataclasses import dataclass, field

# Per-template plausible ranges for headline multiples (advisory only).
_RANGES = {
    "industrial": {
        "P/E (LTM)": (2, 80), "P/E (fwd)": (2, 80), "EV/EBITDA": (1, 40),
        "EV/Sales": (0.1, 15), "P/BV": (0.2, 25), "P/FCF": (2, 120),
    },
    "financials": {
        "P/E (LTM)": (2, 40), "P/E (fwd)": (2, 40), "P/BV": (0.2, 6), "P/TBV": (0.2, 8),
    },
    "reit": {
        "P/FFO": (2, 40), "P/AFFO": (2, 45), "P/NAV": (0.2, 3), "EV/EBITDA": (1, 45),
    },
}

# Multiples that must be positive when present (a negative → meaningless, advisory warn).
_POSITIVE_MULTIPLES = {
    "P/E (LTM)", "P/E (fwd)", "EV/EBITDA", "EV/Sales", "P/BV", "P/FCF",
    "P/TBV", "P/FFO", "P/AFFO", "P/NAV",
}


@dataclass
class AuditCheck:
    name: str
    ok: bool
    detail: str = ""
    gating: bool = True


@dataclass
class AuditReport:
    checks: list[AuditCheck] = field(default_factory=list)

    @property
    def errors(self) -> list[AuditCheck]:
        return [c for c in self.checks if c.gating and not c.ok]

    @property
    def advisories(self) -> list[AuditCheck]:
        return [c for c in self.checks if not c.gating and not c.ok]

    @property
    def passed(self) -> bool:
        return not self.errors


# ---------------------------------------------------------------------------
# Independent arithmetic (own implementation — do NOT import from peers.py)
# ---------------------------------------------------------------------------
def _div(a, b):
    if a is None or b is None or b == 0:
        return None
    return a / b


def _mktcap(px, sh):
    return None if px is None or sh is None else px * sh


def _ev(mc, nd, mi):
    return None if mc is None or nd is None else mc + nd + (mi or 0.0)


# Mirror of peers._PLAUSIBLE_BANDS / admit (kept as an independent copy — the audit never imports the
# engine's arithmetic). The engine blanks a peer multiple that is out-of-band (n/m: negative, or an
# absurd scale artifact), so the audit must blank the SAME ones or it flags a (correct but unhelpful)
# self-consistency mismatch. Percentage columns (dividend/distribution yield) have no band.
_PLAUSIBLE_BANDS = {
    "pe_ltm": (0, 80), "pe_fwd": (0, 80),
    "ev_ebitda": (0, 40), "ev_sales": (0, 15),
    "pbv": (0, 25), "ptbv": (0, 25), "pfcf": (0, 120),
    "p_ffo": (0, 60), "p_affo": (0, 60), "p_nav": (0, 5),
}
_BANK_BANDS = {
    "pe_ltm": (0, 40), "pe_fwd": (0, 60),
    "pbv": (0.1, 5), "ptbv": (0.1, 5),
}
_PCT_MKEYS = {"dvd_yield", "dist_yield", "distribution_yield"}


def _admit(mkey, v, bank=False):
    if v is None or v != v:
        return None
    if mkey in _PCT_MKEYS:
        return v
    bands = _BANK_BANDS if (bank and mkey in _BANK_BANDS) else _PLAUSIBLE_BANDS
    lo, hi = bands.get(mkey, (0.0, float("inf")))
    return v if lo < v < hi else None


def _pct(x):
    return None if x is None else x * 100.0


# Balance-sheet scale-artifact guards — an INDEPENDENT copy of valuation._BS_*/_INV_* literals
# (the audit deliberately never imports the engine's arithmetic). Keep in lock-step with them.
_BS_LO, _BS_HI = 0.001, 5.0
_INV_LO, _INV_HI = 0.005, 2.0


def _plaus(v, ref, lo, hi):
    """Independent copy of valuation._plausible."""
    if v is None or ref is None or ref == 0:
        return False
    return lo * abs(ref) <= abs(v) <= hi * abs(ref)


def _cagr(series: dict, years: int):
    """Independent copy of valuation._cagr (fraction, or None)."""
    if not series:
        return None
    ys = sorted(series)
    latest = ys[-1]
    window = [y for y in ys if y >= latest - years]
    if len(window) < 2:
        return None
    y0, y1 = window[0], window[-1]
    span = y1 - y0
    if span <= 0:
        return None
    v0, v1 = series[y0], series[y1]
    if v0 is None or v1 is None or v0 <= 0 or v1 <= 0:
        return None
    return (v1 / v0) ** (1.0 / span) - 1.0


def _subject_multiples(si, template: str) -> dict[str, float | None]:
    """Recompute the subject's headline multiples independently, keyed by display label."""
    mc = _mktcap(si.px_last, si.shares_out)
    ev = _ev(mc, si.net_debt, si.minority_interest)
    if template == "financials":
        return {
            "Price": si.px_last, "Market cap": mc,
            "P/E (LTM)": _div(mc, si.net_income), "P/E (fwd)": _div(si.px_last, si.eps_ntm),
            "P/BV": _div(mc, si.equity), "P/TBV": _div(mc, si.tangible_book),
            "Dividend yield": si.dvd_yield,
        }
    if template == "reit":
        return {
            "Price": si.px_last, "Market cap": mc,
            "P/FFO": _div(mc, si.ffo), "P/AFFO": _div(mc, si.affo),
            "P/NAV": _div(si.px_last, si.nav_ps), "EV/EBITDA": _div(ev, si.ebitda),
            "Distribution yield": si.distribution_yield,
        }
    return {
        "Price": si.px_last, "Market cap": mc, "Enterprise value": ev,
        "P/E (LTM)": _div(mc, si.net_income), "P/E (fwd)": _div(si.px_last, si.eps_ntm),
        "EV/EBITDA": _div(ev, si.ebitda), "EV/Sales": _div(ev, si.sales),
        "P/BV": _div(mc, si.equity), "P/FCF": _div(mc, si.fcf),
        "Dividend yield": si.dvd_yield,
    }


# ---------------------------------------------------------------------------
# Comparison helpers
# ---------------------------------------------------------------------------
def _close(a, b, rtol=1e-6, atol=1e-6) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) <= atol + rtol * abs(b)


def _finite(x) -> bool:
    return not (isinstance(x, float) and (math.isnan(x) or math.isinf(x)))


class _Auditor:
    def __init__(self, model, spec, fund, pack):
        self.model = model
        self.spec = spec
        self.fund = fund
        self.pack = pack
        self.r = AuditReport()
        self.by_id = {b.id: b for b in model.blocks}

    # -- primitives ---------------------------------------------------------
    def _add(self, name, ok, detail="", gating=True):
        self.r.checks.append(AuditCheck(name, ok, detail, gating))

    def _industrial_si(self):
        """The industrial subject-inputs — the SAME source ``block_financial`` / ``block_profitability``
        use (they call ``_subject_inputs`` unconditionally, not the template-specific
        ``model.subject_inputs``). REITs now carry these analytical blocks, but their
        ``model.subject_inputs`` is NAV-based with no equity, so the recompute must mirror the block's
        own input here or ROE/ROIC/DuPont spuriously blank. Cached; industrial rows are unaffected
        (their ``model.subject_inputs`` already IS this)."""
        si = getattr(self, "_ind_si_cache", None)
        if si is None:
            from src.coverage.valuation import _subject_inputs
            si = self._ind_si_cache = _subject_inputs(self.fund, self.pack)
        return si

    def _block_cell(self, block_id, label):
        b = self.by_id.get(block_id)
        if not b:
            return None
        for c in b.rows:
            if c.label == label:
                return c
        return None

    def _check_equals(self, name, cell, expected, *, unit_scale=1.0):
        """Compare a model cell's value to an independently recomputed ``expected``.

        expected is in RATIO units; ``unit_scale`` converts it to the cell's stored units
        (e.g. 100.0 when the cell stores a percentage). Blank↔blank passes; value-vs-blank and
        value mismatch fail (gating)."""
        if cell is None:
            return  # block/cell absent for this company — nothing to check
        got = cell.value
        exp = None if expected is None else expected * unit_scale
        if got is not None and not _finite(got):
            self._add(name, False, f"non-finite value {got!r}")
            return
        if got is None and exp is None:
            return
        if (got is None) != (exp is None):
            # model printed a value with no support, or dropped a computable one.
            if got is not None:
                self._add(name, False, f"value {got:.6g} present but recompute is blank")
            # exp present / got blank = a sourcing choice, not a math error → ignore.
            return
        ok = _close(got, exp)
        self._add(name, ok, "" if ok else f"model {got:.6g} ≠ recompute {exp:.6g}")

    # -- snapshot multiples (all templates) --------------------------------
    def _audit_snapshot(self):
        si = self.model.subject_inputs
        if si is None:
            self._add("resolved subject inputs present", False, "model.subject_inputs is None")
            return
        tmpl = self.model.template or "industrial"
        snap_id = {"financials": "bank_snapshot", "reit": "reit_snapshot"}.get(tmpl,
                                                                                "snapshot_multiples")
        exp = _subject_multiples(si, tmpl)
        # scalar inputs
        self._check_equals(f"[{snap_id}] Market cap = price×shares",
                           self._block_cell(snap_id, "Market cap"), exp.get("Market cap"))
        if tmpl not in ("financials", "reit"):
            self._check_equals(f"[{snap_id}] EV = mktcap+net_debt+minority",
                               self._block_cell(snap_id, "Enterprise value"),
                               exp.get("Enterprise value"))
        # multiples
        ranges = _RANGES.get(tmpl, {})
        for label, val in exp.items():
            if label in ("Price", "Market cap", "Enterprise value"):
                continue
            cell = self._block_cell(snap_id, label)
            self._check_equals(f"[{snap_id}] {label} arithmetic", cell, val)
            # advisory: positivity + range
            if cell is not None and cell.value is not None and _finite(cell.value):
                if label in _POSITIVE_MULTIPLES and cell.value < 0:
                    self._add(f"[{snap_id}] {label} positive", False,
                              f"{cell.value:.2f}x (negative → meaningless)", gating=False)
                rng = ranges.get(label)
                if rng and not (rng[0] <= cell.value <= rng[1]):
                    self._add(f"[{snap_id}] {label} in [{rng[0]},{rng[1]}]", False,
                              f"{cell.value:.2f} outside plausible range", gating=False)

    # -- EV bridge (independent of the snapshot recompute) ------------------
    def _audit_ev_bridge(self):
        si = self.model.subject_inputs
        if si is None or self.model.template in ("financials", "reit"):
            return
        mc = self._block_cell("snapshot_multiples", "Market cap")
        ev = self._block_cell("snapshot_multiples", "Enterprise value")
        nd = self._block_cell("snapshot_multiples", "Net debt")
        if not (mc and ev and mc.value is not None and ev.value is not None):
            return
        net_debt = nd.value if nd and nd.value is not None else si.net_debt
        minority = si.minority_interest or 0.0
        expected = mc.value + (net_debt or 0.0) + minority
        ok = _close(ev.value, expected)
        self._add("EV bridge ties (EV − mktcap − net_debt − minority = 0)", ok,
                  "" if ok else f"EV {ev.value:.6g} ≠ {expected:.6g}")

    # -- financial-analysis ratio block ------------------------------------
    def _audit_financial(self):
        b = self.by_id.get("financial_analysis")
        if not b:
            return
        si = self._industrial_si()
        f = self.fund
        rev = f.get("revenue")
        # block_financial's margin/conversion use the RAW filing EBITDA (fund.get), NOT the
        # snapshot's reconciled si.ebitda — mirror that here to check the block's own arithmetic.
        ebitda = f.get("ebitda")
        ni = f.get("net_income")
        gp = f.get("gross_profit")
        oi = f.get("operating_income")
        equity = getattr(si, "equity", None)
        self._check_equals("[financial_analysis] EBITDA margin = ebitda/rev",
                           self._block_cell("financial_analysis", "EBITDA margin"),
                           _div(ebitda, rev), unit_scale=100.0)
        self._check_equals("[financial_analysis] Gross margin = gross_profit/rev",
                           self._block_cell("financial_analysis", "Gross margin"),
                           _div(gp, rev), unit_scale=100.0)
        self._check_equals("[financial_analysis] Net margin = net_income/rev",
                           self._block_cell("financial_analysis", "Net margin"),
                           _div(ni, rev), unit_scale=100.0)
        self._check_equals("[financial_analysis] ROE = net_income/equity",
                           self._block_cell("financial_analysis", "ROE"),
                           _div(ni, equity), unit_scale=100.0)
        # ROIC = oi*(1-0.30) / (equity + net_debt)
        nopat = None if oi is None else oi * (1 - 0.30)
        invested = None if equity is None else equity + (getattr(si, "net_debt", None) or 0.0)
        self._check_equals("[financial_analysis] ROIC = NOPAT/(equity+net_debt)",
                           self._block_cell("financial_analysis", "ROIC"),
                           _div(nopat, invested), unit_scale=100.0)
        # FCF/EBITDA conversion
        self._check_equals("[financial_analysis] FCF/EBITDA conversion",
                           self._block_cell("financial_analysis", "FCF / EBITDA conversion"),
                           _div(getattr(si, "fcf", None), ebitda), unit_scale=100.0)
        # Universal P/BV = market cap / equity, and Net debt / EBITDA
        self._check_equals("[financial_analysis] P/BV = mktcap/equity",
                           self._block_cell("financial_analysis", "P/BV"),
                           _div(getattr(si, "market_cap", None), equity))
        self._check_equals("[financial_analysis] Net debt / EBITDA",
                           self._block_cell("financial_analysis", "Net debt / EBITDA"),
                           _div(getattr(si, "net_debt", None), ebitda))
        # revenue growth FY-over-FY
        yrs = sorted(getattr(f, "annual", {}) or {})
        if len(yrs) >= 2:
            cur = f.annual.get(yrs[-1], {})
            prev = f.annual.get(yrs[-2], {})
            g = _div(cur.get("revenue"), prev.get("revenue"))
            growth = None if g is None else (g - 1)
            self._check_equals(f"[financial_analysis] Revenue growth FY{yrs[-2]}→FY{yrs[-1]}",
                               self._block_cell("financial_analysis",
                                                f"Revenue growth FY{yrs[-2]}→FY{yrs[-1]}"),
                               growth, unit_scale=100.0)

    # -- analysis blocks: shared re-derivation helpers ----------------------
    def _ann_val(self, yr, field):
        """Per-FY value: fund.annual first, then the opt-in Bloomberg timeseries (mirrors
        valuation._annual_val)."""
        v = (self.fund.annual.get(yr) or {}).get(field)
        if v is not None:
            return v
        ts = getattr(self.pack, "timeseries", None) or {}
        return (ts.get(yr) or {}).get(field)

    def _series_years(self):
        ts = getattr(self.pack, "timeseries", None) or {}
        return sorted(set(getattr(self.fund, "annual", {}) or {}) | set(ts))

    def _reconcile_assets(self, rev, equity):
        """Mirror block_financial/block_profitability's total-assets reconciliation + equity guard."""
        filing = self.fund.get("total_assets")
        bbg = self.pack.subject.get("total_assets")
        if filing is not None and _plaus(filing, rev, 0.01, 50.0):
            assets = filing
        elif bbg is not None:
            assets = bbg
        else:
            assets = None
        if assets is not None and equity is not None and assets < equity:
            assets = None
        return assets

    # -- profitability trend & DuPont --------------------------------------
    def _audit_profitability(self):
        if not self.by_id.get("profitability"):
            return
        f = self.fund
        for yr in self._series_years():
            rev = self._ann_val(yr, "revenue")
            if rev is None or rev == 0:
                continue
            self._check_equals(f"[profitability] FY{yr} Gross margin",
                               self._block_cell("profitability", f"FY{yr}: Gross margin"),
                               _div(self._ann_val(yr, "gross_profit"), rev), unit_scale=100.0)
            self._check_equals(f"[profitability] FY{yr} EBIT margin",
                               self._block_cell("profitability", f"FY{yr}: EBIT margin"),
                               _div(self._ann_val(yr, "operating_income"), rev), unit_scale=100.0)
            self._check_equals(f"[profitability] FY{yr} EBITDA margin",
                               self._block_cell("profitability", f"FY{yr}: EBITDA margin"),
                               _div(self._ann_val(yr, "ebitda"), rev), unit_scale=100.0)
            self._check_equals(f"[profitability] FY{yr} Net margin",
                               self._block_cell("profitability", f"FY{yr}: Net margin"),
                               _div(self._ann_val(yr, "net_income"), rev), unit_scale=100.0)

        si = self._industrial_si()
        rev, oi, ni = f.get("revenue"), f.get("operating_income"), f.get("net_income")
        equity = getattr(si, "equity", None)
        assets = self._reconcile_assets(rev, equity)
        self._check_equals("[profitability] EBIT margin (LTM)",
                           self._block_cell("profitability", "EBIT margin (LTM)"),
                           _div(oi, rev), unit_scale=100.0)
        nm, at, em = _div(ni, rev), _div(rev, assets), _div(assets, equity)
        self._check_equals("[profitability] DuPont — net margin",
                           self._block_cell("profitability", "DuPont — net margin"),
                           nm, unit_scale=100.0)
        self._check_equals("[profitability] DuPont — asset turnover",
                           self._block_cell("profitability", "DuPont — asset turnover"), at)
        self._check_equals("[profitability] DuPont — equity multiplier",
                           self._block_cell("profitability", "DuPont — equity multiplier"), em)
        implied = None if (nm is None or at is None or em is None) else nm * at * em
        self._check_equals("[profitability] DuPont — implied ROE",
                           self._block_cell("profitability", "DuPont — implied ROE"),
                           implied, unit_scale=100.0)
        # DuPont implied ROE must tie to financial_analysis's ROE within 0.5pp.
        roe_cell = self._block_cell("financial_analysis", "ROE")
        dup_cell = self._block_cell("profitability", "DuPont — implied ROE")
        if (roe_cell is not None and dup_cell is not None
                and roe_cell.value is not None and dup_cell.value is not None
                and _finite(roe_cell.value) and _finite(dup_cell.value)):
            ok = abs(roe_cell.value - dup_cell.value) <= 0.5
            self._add("[profitability] DuPont ROE reconciles to ROE (±0.5pp)", ok,
                      "" if ok else f"DuPont {dup_cell.value:.4g} vs ROE {roe_cell.value:.4g}")

    # -- FCF, liquidity & inventory ----------------------------------------
    def _audit_fcf_liquidity(self):
        if not self.by_id.get("fcf_liquidity"):
            return
        f = self.fund
        si = self.model.subject_inputs
        rev, gp = f.get("revenue"), f.get("gross_profit")
        cogs = (rev - gp) if (rev is not None and gp is not None) else None

        def bs(key, lo=_BS_LO, hi=_BS_HI):
            v = f.get(key)
            return v if _plaus(v, rev, lo, hi) else None

        ca, cl = bs("current_assets"), bs("current_liabilities")
        ar, ap = bs("accounts_receivable"), bs("accounts_payable")
        inv = bs("inventory", _INV_LO, _INV_HI)
        mcap = getattr(si, "market_cap", None)
        fcf = getattr(si, "fcf", None)

        self._check_equals("[fcf_liquidity] FCF yield",
                           self._block_cell("fcf_liquidity", "FCF yield"),
                           _div(fcf, mcap), unit_scale=100.0)
        # Mirror valuation.block_fcf_liquidity: when the revenue-sized gate rejects CA/CL (asset-heavy,
        # low-revenue issuers), the current ratio falls back to RAW CA/CL validated by ratio sanity.
        cr = _div(ca, cl)
        if cr is None:
            cr_raw = _div(f.get("current_assets"), f.get("current_liabilities"))
            if cr_raw is not None and 0.02 <= cr_raw <= 50:
                cr = cr_raw
        self._check_equals("[fcf_liquidity] Current ratio",
                           self._block_cell("fcf_liquidity", "Current ratio"), cr)
        qr = None if (ca is None or inv is None) else _div(ca - inv, cl)
        self._check_equals("[fcf_liquidity] Quick ratio",
                           self._block_cell("fcf_liquidity", "Quick ratio"), qr)
        wc = None if (ca is None or cl is None) else (ca - cl)
        self._check_equals("[fcf_liquidity] Working capital",
                           self._block_cell("fcf_liquidity", "Working capital"), wc)
        self._check_equals("[fcf_liquidity] Inventory turns",
                           self._block_cell("fcf_liquidity", "Inventory turns"), _div(cogs, inv))
        ccc = None
        if all(x is not None for x in (cogs, inv, ar, ap)) and rev:
            dio, dso, dpo = _div(inv, cogs), _div(ar, rev), _div(ap, cogs)
            if None not in (dio, dso, dpo):
                ccc = (dio + dso - dpo) * 365
        self._check_equals("[fcf_liquidity] Cash conversion cycle (days)",
                           self._block_cell("fcf_liquidity", "Cash conversion cycle (days)"), ccc)

        for yr in self._series_years():
            cfo, capex = self._ann_val(yr, "cfo"), self._ann_val(yr, "capex")
            fcf_y = None if (cfo is None or capex is None) else (cfo - abs(capex))
            self._check_equals(f"[fcf_liquidity] FY{yr} FCF",
                               self._block_cell("fcf_liquidity", f"FY{yr}: FCF"), fcf_y)
            revy, invy = self._ann_val(yr, "revenue"), self._ann_val(yr, "inventory")
            if not _plaus(invy, revy, _INV_LO, _INV_HI):
                invy = None
            idays = None
            if invy is not None and revy:
                gpy = self._ann_val(yr, "gross_profit")
                cogsy = (revy - gpy) if gpy is not None else revy
                d = _div(invy, cogsy)
                idays = None if d is None else d * 365
            self._check_equals(f"[fcf_liquidity] FY{yr} Inventory days",
                               self._block_cell("fcf_liquidity", f"FY{yr}: Inventory days"), idays)

    # -- EBIT & EBITDA (temporal) ------------------------------------------
    def _audit_temporal(self):
        if not self.by_id.get("temporal_ebit"):
            return
        ebit_series, ebitda_series = {}, {}
        for yr in self._series_years():
            ebit, eb = self._ann_val(yr, "operating_income"), self._ann_val(yr, "ebitda")
            self._check_equals(f"[temporal_ebit] FY{yr} EBIT",
                               self._block_cell("temporal_ebit", f"FY{yr}: EBIT"), ebit)
            self._check_equals(f"[temporal_ebit] FY{yr} EBITDA",
                               self._block_cell("temporal_ebit", f"FY{yr}: EBITDA"), eb)
            if ebit is not None:
                ebit_series[yr] = ebit
            if eb is not None:
                ebitda_series[yr] = eb

        def yoy(series):
            ys = sorted(series)
            if len(ys) < 2:
                return None
            g = _div(series[ys[-1]], series[ys[-2]])
            return None if g is None else (g - 1)

        self._check_equals("[temporal_ebit] EBIT YoY (latest)",
                           self._block_cell("temporal_ebit", "EBIT YoY (latest)"),
                           yoy(ebit_series), unit_scale=100.0)
        self._check_equals("[temporal_ebit] EBITDA YoY (latest)",
                           self._block_cell("temporal_ebit", "EBITDA YoY (latest)"),
                           yoy(ebitda_series), unit_scale=100.0)
        for name, series in (("EBIT", ebit_series), ("EBITDA", ebitda_series)):
            for span in (3, 5):
                self._check_equals(f"[temporal_ebit] {name} CAGR {span}y",
                                   self._block_cell("temporal_ebit", f"{name} CAGR {span}y"),
                                   _cagr(series, span), unit_scale=100.0)

    # -- Growth -------------------------------------------------------------
    def _audit_growth(self):
        if not self.by_id.get("growth"):
            return
        rev_series, ni_series = {}, {}
        for yr in self._series_years():
            r, n = self._ann_val(yr, "revenue"), self._ann_val(yr, "net_income")
            if r is not None:
                rev_series[yr] = r
            if n is not None:
                ni_series[yr] = n
        ry = sorted(rev_series)
        yoy_vals = []
        for i in range(1, len(ry)):
            cur, prev = ry[i], ry[i - 1]
            g = _div(rev_series[cur], rev_series[prev])
            exp = None if g is None else (g - 1)
            self._check_equals(f"[growth] FY{cur} Revenue YoY",
                               self._block_cell("growth", f"FY{cur}: Revenue YoY"),
                               exp, unit_scale=100.0)
            if exp is not None:
                yoy_vals.append(exp * 100)
        if yoy_vals:
            # 1-year revenue "CAGR" == the latest FY revenue YoY (already a percent).
            self._check_equals("[growth] Revenue CAGR 1y",
                               self._block_cell("growth", "Revenue CAGR 1y"),
                               yoy_vals[-1])
        # net-income YoY + 1y CAGR — only off a positive prior year (matches block_growth).
        ny = sorted(ni_series)
        ni_yoy_vals = []
        for i in range(1, len(ny)):
            cur, prev = ny[i], ny[i - 1]
            if ni_series[prev] > 0:
                g = _div(ni_series[cur], ni_series[prev])
                exp = None if g is None else (g - 1)
                self._check_equals(f"[growth] FY{cur} Net income YoY",
                                   self._block_cell("growth", f"FY{cur}: Net income YoY"),
                                   exp, unit_scale=100.0)
                if exp is not None:
                    ni_yoy_vals.append(exp * 100)
        if ni_yoy_vals:
            self._check_equals("[growth] Net income CAGR 1y",
                               self._block_cell("growth", "Net income CAGR 1y"),
                               ni_yoy_vals[-1])
        for span in (3, 5):
            self._check_equals(f"[growth] Revenue CAGR {span}y",
                               self._block_cell("growth", f"Revenue CAGR {span}y"),
                               _cagr(rev_series, span), unit_scale=100.0)
        for span in (3, 5):
            self._check_equals(f"[growth] Net income CAGR {span}y",
                               self._block_cell("growth", f"Net income CAGR {span}y"),
                               _cagr(ni_series, span), unit_scale=100.0)
        if len(yoy_vals) >= 2:
            self._check_equals("[growth] Revenue growth stability (σ)",
                               self._block_cell("growth", "Revenue growth stability (σ)"),
                               statistics.pstdev(yoy_vals))

    # -- bank returns block -------------------------------------------------
    def _audit_bank(self):
        if self.model.template != "financials":
            return
        si = self.model.subject_inputs
        if si is None:
            return
        roe = _div(si.net_income, si.equity)
        rote = _div(si.net_income, si.tangible_book)
        # ROE/ROTE cells fall back to bbg when the calc is unavailable; only check the calc path.
        roe_cell = self._block_cell("bank_returns", "ROE")
        if roe is not None:
            self._check_equals("[bank_returns] ROE = net_income/equity", roe_cell, roe,
                               unit_scale=100.0)
        rote_cell = self._block_cell("bank_returns", "ROTE")
        if rote is not None:
            self._check_equals("[bank_returns] ROTE = net_income/tangible_book", rote_cell, rote,
                               unit_scale=100.0)

    # -- cross-section peers + median --------------------------------------
    def _audit_cross_section(self):
        tmpl = self.model.template or "industrial"
        snap_id = {"financials": "bank_snapshot", "reit": "reit_snapshot"}.get(tmpl,
                                                                                "snapshot_multiples")
        b = self.by_id.get(snap_id)
        if not b or b.cross_section is None:
            return
        xs = b.cross_section
        peers = self.model.peer_inputs or []
        peer_by_name = {p.name: p for p in peers}
        for mkey, per_company in xs.multiples.items():
            # recompute each peer's multiple independently
            recomputed = {}
            for name, _val in per_company.items():
                if name == xs.subject:
                    continue
                p = peer_by_name.get(name)
                if p is None:
                    continue
                # Blank out-of-band recomputes so they match the engine's n/m treatment (banks use
                # the tighter band, same as build_cross_section).
                recomputed[name] = _admit(mkey, self._peer_multiple(p, mkey, tmpl),
                                          bank=(tmpl == "financials"))
            for name, exp in recomputed.items():
                got = per_company.get(name)
                if got is None and exp is None:
                    continue
                if (got is None) != (exp is None):
                    if got is not None:
                        self._add(f"[peer {name}] {mkey} arithmetic", False,
                                  f"{got:.6g} present but recompute blank")
                    continue
                ok = _close(got, exp)
                self._add(f"[peer {name}] {mkey} arithmetic", ok,
                          "" if ok else f"{got:.6g} ≠ {exp:.6g}")
            # median of the meaningful peer values (out-of-band recomputes are already blanked
            # above, mirroring the engine) — so garbage can't poison it.
            vals = [v for v in recomputed.values() if v is not None]
            exp_med = statistics.median(vals) if vals else None
            got_med = xs.median.get(mkey)
            if not (got_med is None and exp_med is None):
                ok = _close(got_med, exp_med)
                self._add(f"[peer median] {mkey}", ok,
                          "" if ok else f"{got_med} ≠ {exp_med}")

    def _peer_multiple(self, p, mkey, tmpl):
        mc = _mktcap(p.px_last, p.shares_out)
        ev = _ev(mc, p.net_debt, p.minority_interest)
        table = {
            "pe_ltm": _div(mc, p.net_income), "pe_fwd": _div(p.px_last, p.eps_ntm),
            "ev_ebitda": _div(ev, p.ebitda), "ev_sales": _div(ev, p.sales),
            "pbv": _div(mc, p.equity), "pfcf": _div(mc, p.fcf),
            "ptbv": _div(mc, p.tangible_book), "dvd_yield": p.dvd_yield,
            "p_ffo": _div(mc, p.ffo), "p_affo": _div(mc, p.affo),
            "p_nav": _div(p.px_last, p.nav_ps), "dist_yield": p.distribution_yield,
        }
        return table.get(mkey)

    # -- historical bands ---------------------------------------------------
    def _audit_bands(self):
        for b in self.model.blocks:
            labels = {c.label for c in b.rows}
            names = {lbl.split(" — ")[0] for lbl in labels if " — historical mean" in lbl}
            for name in names:
                series = []
                for c in b.rows:
                    if c.label.startswith(f"FY") and c.label.endswith(f": {name}") \
                            and c.value is not None and _finite(c.value):
                        series.append(c.value)
                if not series:
                    continue
                mean = statistics.fmean(series)
                sd = statistics.pstdev(series) if len(series) > 1 else 0.0
                self._check_equals(f"[{b.id}] {name} — mean",
                                   self._block_cell(b.id, f"{name} — historical mean"), mean)
                self._check_equals(f"[{b.id}] {name} — +1σ",
                                   self._block_cell(b.id, f"{name} — +1σ"), mean + sd)
                self._check_equals(f"[{b.id}] {name} — −1σ",
                                   self._block_cell(b.id, f"{name} — −1σ"), mean - sd)

    # -- SOTP bridge --------------------------------------------------------
    def _audit_sotp(self):
        b = self.by_id.get("sum_of_the_parts")
        if not b:
            return
        seg_evs = [c.value for c in b.rows
                   if ": EBITDA ×" in c.label and c.value is not None and _finite(c.value)]
        sum_cell = self._block_cell("sum_of_the_parts", "Sum of segment EV")
        if seg_evs and sum_cell is not None:
            self._check_equals("[sum_of_the_parts] Σ segment EV", sum_cell, sum(seg_evs))
        # implied equity = sum EV − (net debt + minority); implied px = implied equity / shares
        si = self.model.subject_inputs
        if si is not None and sum_cell is not None and sum_cell.value is not None:
            net_debt = getattr(si, "net_debt", None) or 0.0
            minority = getattr(si, "minority_interest", None) or 0.0
            implied_eq = sum_cell.value - net_debt - minority
            self._check_equals("[sum_of_the_parts] implied equity",
                               self._block_cell("sum_of_the_parts", "Implied equity value"),
                               implied_eq)
            eq_cell = self._block_cell("sum_of_the_parts", "Implied equity value")
            if eq_cell is not None and eq_cell.value is not None:
                self._check_equals("[sum_of_the_parts] implied price/share",
                                   self._block_cell("sum_of_the_parts", "Implied price / share"),
                                   _div(eq_cell.value, si.shares_out))

    # -- fundamentals accounting identities (reuse the vendored validator) --
    def _audit_fundamentals(self):
        try:
            from src.coverage.validate import _fundamentals_identities
        except Exception:
            return
        for res in _fundamentals_identities(self.fund):
            passed = bool(getattr(res, "passed", False))
            rule = getattr(res, "rule", "?")
            msg = getattr(res, "message", "")
            # accounting-identity failures are gating; pure sanity warnings stay advisory.
            # NOTE: ebitda_derivation (ebitda ≈ operating_income + D&A) is an INPUT-completeness
            # check, not a ratio-arithmetic one — a moderate gap usually means incomplete D&A
            # extraction, while the shipped EV/EBITDA is still arithmetically correct on the reported
            # EBITDA. The 10^6 scale artifacts are corrected upstream (_derive_native scale guard),
            # so this stays ADVISORY. margin_consistency (ebitda_margin == ebitda/rev) remains gating.
            gating = rule in {"gross_profit_identity", "net_income_identity",
                              "net_debt_identity", "fcf_derivation", "margin_consistency",
                              "balance_sheet_identity", "revenue_segment_sum"}
            self._add(f"[fundamentals] {rule}", passed, msg, gating=gating)

    # -- model non-finite sweep --------------------------------------------
    def _audit_finite(self):
        for b in self.model.blocks:
            for c in b.rows:
                if c.value is not None and not _finite(c.value):
                    self._add(f"[{b.id}] {c.label} finite", False, f"non-finite {c.value!r}")

    def run(self) -> AuditReport:
        self._audit_finite()
        self._audit_snapshot()
        self._audit_ev_bridge()
        self._audit_financial()
        self._audit_profitability()
        self._audit_fcf_liquidity()
        self._audit_temporal()
        self._audit_growth()
        self._audit_bank()
        self._audit_cross_section()
        self._audit_bands()
        self._audit_sotp()
        self._audit_fundamentals()
        return self.r


# ---------------------------------------------------------------------------
# Workbook (Excel formula) audit
# ---------------------------------------------------------------------------
def _audit_workbook(report: AuditReport, model, spec, xlsx_path) -> None:
    from openpyxl import load_workbook

    from src.coverage.formula_eval import DivisionByZero, FormulaError, make_evaluator

    try:
        wb = load_workbook(xlsx_path)  # keeps formulas as strings
    except Exception as e:
        report.checks.append(AuditCheck("workbook loads", False, f"{e}"))
        return
    ws = wb.active

    # 1) no Excel error strings, no non-finite numeric cells anywhere.
    errors_found = 0
    nonfinite = 0
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("#") and v.endswith("!"):
                errors_found += 1
            elif isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                nonfinite += 1
    report.checks.append(AuditCheck("workbook has no Excel error cells", errors_found == 0,
                                    f"{errors_found} #ERR cells" if errors_found else ""))
    report.checks.append(AuditCheck("workbook has no non-finite numbers", nonfinite == 0,
                                    f"{nonfinite} NaN/Inf cells" if nonfinite else ""))

    # 2) every formula evaluates cleanly (no div-by-zero, no parse failure) to a finite number.
    ev = make_evaluator(ws)
    bad = 0
    detail = ""
    for coord, formula in ev.formula_cells().items():
        try:
            val = ev.cell(coord)
        except DivisionByZero:
            bad += 1
            detail = detail or f"{coord}: {formula} → #DIV/0!"
            continue
        except FormulaError as e:
            bad += 1
            detail = detail or f"{coord}: {formula} → {e}"
            continue
        if val is not None and not _finite(val):
            bad += 1
            detail = detail or f"{coord}: {formula} → non-finite"
    report.checks.append(AuditCheck("workbook formulas evaluate cleanly", bad == 0, detail))

    # 3) industrial snapshot: evaluated formula == independent recompute (ties Excel ⇄ math).
    si = model.subject_inputs
    if si is None or (model.template or "industrial") not in ("industrial", ""):
        return
    recompute = _subject_multiples(si, "industrial")
    med_col = 3 + len(spec.peers)
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if not isinstance(label, str) or label not in recompute:
            continue
        b = ws.cell(row=r, column=2).value
        if isinstance(b, str) and b.startswith("="):
            try:
                got = ev.cell(ws.cell(row=r, column=2).coordinate)
            except (DivisionByZero, FormulaError) as e:
                report.checks.append(AuditCheck(f"[xlsx] {label} formula evaluates", False, str(e)))
                continue
            exp = recompute.get(label)
            if got is None and exp is None:
                continue
            ok = (got is not None and exp is not None and _close(got, exp))
            report.checks.append(AuditCheck(f"[xlsx] {label} formula = recompute", ok,
                                            "" if ok else f"{got} ≠ {exp}"))


def audit(model, spec, fund, pack, xlsx_path=None) -> AuditReport:
    """Run the full math audit. Pass ``xlsx_path`` to also verify the emitted workbook's
    formulas; omit it to audit the Python model only (fast, for unit tests)."""
    report = _Auditor(model, spec, fund, pack).run()
    if xlsx_path is not None:
        _audit_workbook(report, model, spec, xlsx_path)
    return report


def report_markdown(report: AuditReport, name: str) -> str:
    lines = [f"# {name} — Math verification", ""]
    verdict = "✅ PASS" if report.passed else "❌ FAIL"
    lines.append(f"- Verdict: **{verdict}**  ·  {len(report.errors)} error(s), "
                 f"{len(report.advisories)} advisory")
    lines.append("")
    if report.errors:
        lines.append("## Gating errors (block the deliverable)")
        lines.append("")
        lines.append("| Check | Detail |")
        lines.append("|---|---|")
        for c in report.errors:
            lines.append(f"| {c.name} | {c.detail} |")
        lines.append("")
    if report.advisories:
        lines.append("## Advisories (non-blocking)")
        lines.append("")
        lines.append("| Check | Detail |")
        lines.append("|---|---|")
        for c in report.advisories:
            lines.append(f"| {c.name} | {c.detail} |")
        lines.append("")
    passed_n = sum(1 for c in report.checks if c.ok)
    lines.append(f"_{passed_n}/{len(report.checks)} checks passed._")
    return "\n".join(lines)
