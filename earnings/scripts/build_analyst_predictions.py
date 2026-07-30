#!/usr/bin/env python3
"""build_analyst_predictions.py — analyst directional calls from the frozen
Analyst Expectation workbook.

Parses the expectations block (ticker rows x quarter columns 2Q24-2Q26),
normalizes labels through configs/analyst_labels.yaml (semantics are data, not
code), and maps tickers to slugs with the same chain as the report calendar.

Output: outputs/analyst_predictions.parquet
  ticker, slug, symbol, period, label_raw, direction (+1/0/-1),
  strength (1.0 full / 0.5 soft / 0.0 neutral), source_row
"""
from __future__ import annotations

import sys
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from earnlib import bootstrap as bs

import pandas as pd
import yaml


def block_period(block: str) -> str:
    """'3Q24' -> '2024-3T'."""
    return f"20{block[2:]}-{block[0]}T"


def main() -> None:
    cal_cfg = yaml.safe_load(open(bs.EARNINGS_ROOT / "configs" / "analyst_calendar.yaml"))
    lab_cfg = yaml.safe_load(open(bs.EARNINGS_ROOT / "configs" / "analyst_labels.yaml"))
    labels = lab_cfg["labels"]

    xlsx = bs.EARNINGS_ROOT / cal_cfg["source"]
    import hashlib

    assert hashlib.sha256(xlsx.read_bytes()).hexdigest() == cal_cfg["source_sha256"], \
        "frozen workbook changed on disk"

    import openpyxl

    warnings.filterwarnings("ignore", module="openpyxl")
    ws = openpyxl.load_workbook(xlsx, data_only=True)[cal_cfg["sheet"]]

    universe = bs.load_universe(bs.load_config(), "full")
    tick2slug = {v["ticker"].upper(): slug for slug, v in universe.items()}
    tick2slug.update({k.upper(): v for k, v in (cal_cfg.get("ticker_overrides") or {}).items()})
    slug2sym = {slug: v["symbol"] for slug, v in universe.items()}
    unmappable = {t.upper() for t in cal_cfg.get("unmappable_tickers") or ()}

    rows, unknown, unmapped = [], [], set()
    for r in range(2, 61):
        cell = ws.cell(row=r, column=2).value
        if not cell:
            continue
        ticker = str(cell).strip()
        if ticker != ticker.upper():        # sector header rows are mixed-case
            continue
        for block, col in lab_cfg["quarter_columns"].items():
            v = ws.cell(row=r, column=col).value
            if v is None or str(v).strip() == "":
                continue
            raw = str(v).strip()
            if raw not in labels:
                unknown.append((ticker, block, raw))
                continue
            slug = tick2slug.get(ticker.upper())
            if slug is None and ticker.upper() not in unmappable:
                unmapped.add(ticker)
            spec = labels[raw]
            rows.append({
                "ticker": ticker, "slug": slug,
                "symbol": slug2sym.get(slug) if slug else None,
                "period": block_period(block), "label_raw": raw,
                "direction": int(spec["direction"]),
                "strength": float(spec["strength"]), "source_row": r,
            })

    if lab_cfg.get("strict") and unknown:
        for u in unknown:
            print("UNKNOWN LABEL:", u)
        raise SystemExit(f"{len(unknown)} unknown labels")
    assert not unmapped, f"unmapped prediction tickers: {sorted(unmapped)}"

    df = pd.DataFrame(rows)
    df.to_parquet(bs.OUTPUTS_DIR / "analyst_predictions.parquet", index=False)

    # ---- verify block ----
    print(f"predictions: {len(df)} across {df['ticker'].nunique()} tickers, "
          f"{df['period'].nunique()} quarters")
    print("label counts:", df["label_raw"].value_counts().to_dict())
    print("directional (|dir|=1):", int((df["direction"] != 0).sum()),
          "| full-strength:", int((df["strength"] == 1.0).sum()),
          "| soft:", int((df["strength"] == 0.5).sum()),
          "| neutral:", int((df["direction"] == 0).sum()))
    print("mapped to slugs:", int(df["slug"].notna().sum()),
          "| unmappable:", sorted(df.loc[df["slug"].isna(), "ticker"].unique()))
    print("\nper quarter:")
    print(df.groupby("period")["ticker"].count().to_string())


if __name__ == "__main__":
    main()
