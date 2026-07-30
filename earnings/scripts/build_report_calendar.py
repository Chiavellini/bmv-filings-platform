#!/usr/bin/env python3
"""build_report_calendar.py — human-curated report calendar from the frozen
Analyst Expectation workbook.

Parses the eight "Calendario <QQYY>" blocks (ticker, date, Cierre/Apertura),
applies the pre-committed year corrections for the 2Q25/3Q25 entry bugs, maps
Excel tickers to universe slugs, and verifies every corrected date against
period-end bounds and (where available) the XBRL filing date.

Output: outputs/report_calendar.parquet
  ticker, slug, symbol, period (YYYY-NT), cal_date (date), session
  (after_close | pre_open), year_corrected (bool), source_block, source_row
"""
from __future__ import annotations

import sys
import warnings
from datetime import datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from earnlib import bootstrap as bs

import pandas as pd
import yaml

_QEND = {"1": "-03-31", "2": "-06-30", "3": "-09-30", "4": "-12-31"}


def block_period(block: str) -> str:
    """'3Q24' -> '2024-3T'."""
    return f"20{block[2:]}-{block[0]}T"


def main() -> None:
    cfg = yaml.safe_load(open(bs.EARNINGS_ROOT / "configs" / "analyst_calendar.yaml"))
    xlsx = bs.EARNINGS_ROOT / cfg["source"]
    import hashlib

    sha = hashlib.sha256(xlsx.read_bytes()).hexdigest()
    assert sha == cfg["source_sha256"], "frozen workbook changed on disk"

    import openpyxl

    warnings.filterwarnings("ignore", module="openpyxl")
    ws = openpyxl.load_workbook(xlsx, data_only=True)[cfg["sheet"]]

    corrections = {c["block"]: c for c in cfg["year_corrections"]}
    session_map = cfg["session_map"]

    universe = bs.load_universe(bs.load_config(), "full")
    tick2slug = {v["ticker"].upper(): slug for slug, v in universe.items()}
    tick2slug.update({k.upper(): v for k, v in (cfg.get("ticker_overrides") or {}).items()})
    slug2sym = {slug: v["symbol"] for slug, v in universe.items()}
    unmappable = {t.upper() for t in cfg.get("unmappable_tickers") or ()}

    rows, unmapped = [], set()
    for block, col in cfg["calendar_columns"].items():
        header = ws.cell(row=1, column=col).value
        assert header == f"Calendario {block}", (block, header)
        period = block_period(block)
        corr = corrections.get(block)
        for r in range(2, 61):
            ticker = ws.cell(row=r, column=col).value
            date_v = ws.cell(row=r, column=col + 1).value
            sess_v = ws.cell(row=r, column=col + 2).value
            if not ticker or not isinstance(date_v, datetime):
                continue
            ticker = str(ticker).strip().upper()
            session = session_map.get(str(sess_v).strip() if sess_v else "")
            assert session is not None, f"unknown session {sess_v!r} at {block} row {r}"
            corrected = False
            cal_date = date_v.date()
            if corr and cal_date.year == corr["wrong_year"]:
                cal_date = cal_date.replace(year=corr["correct_year"])
                corrected = True
            slug = tick2slug.get(ticker)
            if slug is None and ticker not in unmappable:
                unmapped.add(ticker)
            rows.append({
                "ticker": ticker, "slug": slug,
                "symbol": slug2sym.get(slug) if slug else None,
                "period": period, "cal_date": cal_date, "session": session,
                "year_corrected": corrected, "source_block": block, "source_row": r,
            })

    df = pd.DataFrame(rows)
    assert not unmapped, f"unmapped calendar tickers (add to ticker_overrides): {sorted(unmapped)}"

    # -- plausibility verification ------------------------------------------- #
    # HARD bound: the (possibly year-corrected) date must sit in the report
    # window after period end — a wrong year correction lands ~365d off and
    # fails this. The distance to the XBRL filing is a DIAGNOSTIC only: a big
    # gap is exactly the XBRL-lags-the-release phenomenon W2 audits (e.g. CADU
    # filed its 2025-2T XBRL in December), never proof the calendar is wrong.
    bounds = cfg["correction_bounds"]
    events = pd.read_parquet(bs.art_path("events"))
    filed_by_key = {(r["slug"], r["period"]): r["filed_dt"]
                    for _, r in events.iterrows() if r["slug"]}
    viol, lag_days = [], []
    for _, r in df.iterrows():
        pend = datetime.fromisoformat(r["period"][:4] + _QEND[r["period"][5]]).date()
        if not (pend < r["cal_date"] <= pend + timedelta(days=bounds["max_days_after_period_end"])):
            viol.append((r["ticker"], r["period"], str(r["cal_date"]), "outside period window"))
        filed = filed_by_key.get((r["slug"], r["period"]))
        lag_days.append((filed.date() - r["cal_date"]).days if filed is not None else None)
    df["xbrl_minus_cal_days"] = lag_days
    if viol:
        for v in viol:
            print("VIOLATION:", v)
        raise SystemExit(f"{len(viol)} calendar rows failed the period-window bound")
    big_lag = df[df["year_corrected"] & (df["xbrl_minus_cal_days"].abs() > bounds["max_days_vs_xbrl"])]
    if len(big_lag):
        print(f"note: {len(big_lag)} year-corrected rows have |xbrl - calendar| > "
              f"{bounds['max_days_vs_xbrl']}d (late XBRL filers, diagnostic):")
        print(big_lag[["ticker", "period", "cal_date", "xbrl_minus_cal_days"]]
              .to_string(index=False))

    out = bs.OUTPUTS_DIR / "report_calendar.parquet"
    df.to_parquet(out, index=False)

    # ---- verify block ----
    print(f"calendar rows: {len(df)} across {df['ticker'].nunique()} tickers, "
          f"{df['period'].nunique()} quarters")
    print(f"mapped to slugs: {df['slug'].notna().sum()} "
          f"(unmappable by config: {sorted(df.loc[df['slug'].isna(), 'ticker'].unique())})")
    print(f"year-corrected rows: {int(df['year_corrected'].sum())} "
          f"(blocks {sorted(df.loc[df['year_corrected'], 'source_block'].unique())})")
    print(df.groupby("session")["ticker"].count().to_string())
    print("\nper-block counts:")
    print(df.groupby("source_block")["ticker"].count().to_string())
    xb = df["slug"].map(lambda s: True if s else False)
    both = [(r["slug"], r["period"]) in filed_by_key
            for _, r in df.iterrows()]
    print(f"\nrows with a matching XBRL filing event: {sum(both)}")


if __name__ == "__main__":
    main()
